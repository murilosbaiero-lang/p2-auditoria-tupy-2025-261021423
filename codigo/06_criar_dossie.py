#!/usr/bin/env python3
"""Gera o dossiê acadêmico consolidado em DOCX.

O documento combina fontes públicas, planejamento de auditoria e simulações
didáticas. Ele não constitui auditoria independente nem opinião sobre a Tupy.
"""

from __future__ import annotations

import csv
import json
import re
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dossie" / "Dossie_Auditoria_Tupy_2025_261021423.docx"
FIG = ROOT / "tmp" / "figures"
FIG.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE = "F3F6F9"
GRAY = "D9E2F3"
DARK_GRAY = "595959"
RED = "9C0006"
GREEN = "548235"
AMBER = "BF8F00"
WHITE = "FFFFFF"
BLACK = "000000"


def csv_rows(rel: str) -> list[dict[str, str]]:
    with (ROOT / rel).open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def brl(value: float | int, decimals: int = 0) -> str:
    text = f"{abs(value):,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"(R$ {text})" if value < 0 else f"R$ {text}"


def pct(value: float, decimals: int = 1) -> str:
    text = f"{abs(value) * 100:.{decimals}f}".replace(".", ",") + "%"
    return f"({text})" if value < 0 else text


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_together = value


def add_field(paragraph, field_code: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Atualize o campo no Word, se necessário"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    run.append(begin)
    run.append(instr)
    run.append(sep)
    run.append(text)
    run.append(end)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Página ")
    add_field(paragraph, "PAGE")
    paragraph.add_run(" de ")
    add_field(paragraph, "NUMPAGES")


def set_update_fields_on_open(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def set_paragraph_border(paragraph, color: str = NAVY, size: int = 18, position: str = "bottom") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    edge = OxmlElement(f"w:{position}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "1")
    edge.set(qn("w:color"), color)
    borders.append(edge)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_run_font(run, name: str = "Aptos", size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text: str, size: float | None = None, color: str | None = None) -> None:
    """Adiciona marcações simples de Markdown (negrito, itálico e código)."""
    token = re.compile(r"(\*\*.*?\*\*|`.*?`|(?<!\*)\*[^*]+?\*(?!\*))")
    pos = 0
    for match in token.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, color=color)
        raw = match.group(0)
        if raw.startswith("**"):
            run = paragraph.add_run(raw[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif raw.startswith("`"):
            run = paragraph.add_run(raw[1:-1])
            set_run_font(run, name="Aptos Mono", size=(size or 9) - 0.5, color=color)
            run.font.highlight_color = None
        else:
            run = paragraph.add_run(raw[1:-1])
            set_run_font(run, size=size, color=color, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.70)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.30)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in (
        ("Title", 27, NAVY, 0, 10),
        ("Subtitle", 14, DARK_GRAY, 0, 8),
        ("Heading 1", 17, NAVY, 16, 8),
        ("Heading 2", 13.5, BLUE, 12, 5),
        ("Heading 3", 11, DARK_GRAY, 8, 3),
        ("Heading 4", 9.5, DARK_GRAY, 6, 2),
    ):
        st = styles[style_name]
        st.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), st.font.name)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.widow_control = True

    styles["Heading 1"].paragraph_format.page_break_before = False
    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(9.5)
    styles["List Number"].font.name = "Aptos"
    styles["List Number"].font.size = Pt(9.5)

    if "Legenda" not in styles:
        legend = styles.add_style("Legenda", WD_STYLE_TYPE.PARAGRAPH)
    else:
        legend = styles["Legenda"]
    legend.font.name = "Aptos"
    legend.font.size = Pt(8)
    legend.font.italic = True
    legend.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    legend.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    legend.paragraph_format.space_after = Pt(7)

    if "Nota técnica" not in styles:
        note = styles.add_style("Nota técnica", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = styles["Nota técnica"]
    note.font.name = "Aptos"
    note.font.size = Pt(9)
    note.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    note.paragraph_format.left_indent = Inches(0.15)
    note.paragraph_format.right_indent = Inches(0.15)
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(6)
    note.paragraph_format.keep_together = True

    # Marcadores de atualização automática e compatibilidade.
    set_update_fields_on_open(doc)
    compat = doc.settings._element.find(qn("w:compat"))
    if compat is not None:
        setting = OxmlElement("w:compatSetting")
        setting.set(qn("w:name"), "compatibilityMode")
        setting.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
        setting.set(qn("w:val"), "15")
        compat.append(setting)

    core = doc.core_properties
    core.title = "Dossiê Acadêmico de Auditoria — Tupy S.A. — 2025"
    core.subject = "Planejamento e documentação de auditoria conforme NBC TA 200 a 600"
    core.author = "Murilo Sena Baiero"
    core.keywords = "auditoria; Tupy; NBC TA; imobilizado; base sintética; 2025"
    core.comments = "Trabalho acadêmico; não constitui auditoria ou opinião independente."


def set_headers_footers(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run("DOSSIÊ ACADÊMICO DE AUDITORIA  |  TUPY S.A.  |  2025")
        set_run_font(run, size=7.5, color=DARK_GRAY, bold=True)
        set_paragraph_border(p, color=LIGHT_BLUE, size=8, position="bottom")

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Uso acadêmico — não constitui auditoria nem opinião independente  •  ")
        set_run_font(run, size=7.5, color=DARK_GRAY)
        add_page_number(p)
        for run in p.runs:
            set_run_font(run, size=7.5, color=DARK_GRAY)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("UNIVERSIDADE DE BRASÍLIA")
    set_run_font(run, size=10.5, color=DARK_GRAY, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(60)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("DOSSIÊ ACADÊMICO\nDE AUDITORIA")
    set_run_font(run, name="Aptos Display", size=29, color=NAVY, bold=True)
    set_paragraph_border(p, color=NAVY, size=24, position="bottom")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("TUPY S.A. E CONTROLADAS")
    set_run_font(run, name="Aptos Display", size=18, color=BLUE, bold=True)
    p.add_run("\n")
    run = p.add_run("Exercício findo em 31 de dezembro de 2025")
    set_run_font(run, size=12.5, color=DARK_GRAY)

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(4.95)
    data = [
        ("Responsável", "Murilo Sena Baiero"),
        ("Matrícula", "261021423"),
        ("Natureza", "Trabalho individual — P2"),
        ("Referencial", "Estrutura Conceitual e NBC TA 200 a 600"),
        ("Elaboração", "24 de outubro de 2026"),
    ]
    for i, (label, value) in enumerate(data):
        a, b = table.rows[i].cells
        set_cell_shading(a, NAVY)
        set_cell_shading(b, PALE if i % 2 == 0 else WHITE)
        for cell in (a, b):
            set_cell_margins(cell, 110, 120, 110, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        ra = a.paragraphs[0].add_run(label.upper())
        set_run_font(ra, size=8.2, color=WHITE, bold=True)
        rb = b.paragraphs[0].add_run(value)
        set_run_font(rb, size=9.5, color=BLACK, bold=(i < 2))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AVISO DE NATUREZA ACADÊMICA")
    set_run_font(run, size=9.5, color=RED, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.55)
    p.paragraph_format.right_indent = Inches(0.55)
    add_inline(
        p,
        "Este dossiê usa exclusivamente informações públicas e dados artificiais. Não houve acesso aos sistemas, documentos internos ou responsáveis da Tupy. Nenhuma exceção simulada representa constatação sobre a companhia, e nenhuma opinião independente é emitida.",
        size=8.8,
        color=DARK_GRAY,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BRASÍLIA  •  2026")
    set_run_font(run, size=9, color=DARK_GRAY, bold=True)
    doc.add_page_break()


def add_section_title(doc: Document, title: str, subtitle: str | None = None,
                      page_break: bool = False, bookmark: str | None = None, bookmark_id: int = 1):
    if page_break and len(doc.paragraphs) > 0:
        doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0 if page_break else 14)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(title)
    set_run_font(run, name="Aptos Display", size=17, color=NAVY, bold=True)
    set_paragraph_border(p, color=NAVY, size=18, position="bottom")
    if bookmark:
        add_bookmark(p, bookmark, bookmark_id)
    if subtitle:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(8)
        sr = sp.add_run(subtitle)
        set_run_font(sr, size=9.2, color=DARK_GRAY, italic=True)
    return p


def add_kpi_table(doc: Document) -> None:
    rows = [
        ("Ativo total", "R$ 9,392 bi", "DFP 2025"),
        ("Receita", "R$ 9,693 bi", "queda de 9,1%"),
        ("Resultado líquido", "R$ (654,6) mi", "prejuízo no exercício"),
        ("Caixa operacional", "R$ 915,1 mi", "geração positiva"),
        ("Imobilizado líquido", "R$ 2,525 bi", "26,9% do ativo"),
        ("Materialidade global", "R$ 46,96 mi", "0,5% do ativo"),
    ]
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.13), Inches(2.13), Inches(2.13)]
    for j, width in enumerate(widths):
        table.columns[j].width = width
    for r_idx in range(2):
        for c_idx in range(3):
            label, value, context = rows[r_idx * 3 + c_idx]
            cell = table.cell(r_idx, c_idx)
            set_cell_shading(cell, LIGHT_BLUE if r_idx == 0 else PALE)
            set_cell_margins(cell, 130, 130, 130, 130)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(label.upper())
            set_run_font(r, size=7.5, color=NAVY, bold=True)
            p.add_run("\n")
            r = p.add_run(value)
            set_run_font(r, size=14, color=NAVY if "(" not in value else RED, bold=True)
            p.add_run("\n")
            r = p.add_run(context)
            set_run_font(r, size=7.5, color=DARK_GRAY)
    doc.add_paragraph()


def add_dataframe_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None,
                        font_size: float = 7.8, header_color: str = NAVY, alternate: bool = True) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Inches(width)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    hdr.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    hdr.height = Pt(24)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_color)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=WHITE, bold=True)
    for r_idx, values in enumerate(rows):
        cells = table.add_row().cells
        if alternate and r_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, PALE)
        for c_idx, value in enumerate(values):
            cell = cells[c_idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c_idx > 0 and re.fullmatch(r"[\d.(),%xR$\s-]+", str(value)):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_inline(p, str(value), size=font_size, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_note(doc: Document, title: str, text: str, color: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    set_cell_margins(cell, 120, 140, 120, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{title.upper()}  ")
    set_run_font(r, size=8.5, color=NAVY if color != "FCE8E6" else RED, bold=True)
    add_inline(p, text, size=8.8, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def create_figures() -> dict[str, Path]:
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "axes.titlesize": 12,
        "axes.labelsize": 9,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
    })

    # Desempenho comparativo.
    labels = ["Receita", "Lucro\nbruto", "Resultado\noperacional", "Resultado\nlíquido", "Caixa\noperacional"]
    fy24 = np.array([10665.110, 1926.591, 559.667, 82.440, 1353.491])
    fy25 = np.array([9692.948, 1223.151, -156.481, -654.552, 915.137])
    x = np.arange(len(labels))
    fig, ax = plt.subplots(figsize=(9.3, 4.0), dpi=180)
    width = 0.34
    ax.bar(x - width / 2, fy24, width, label="2024", color="#176B87")
    ax.bar(x + width / 2, fy25, width, label="2025", color="#ED7D31")
    ax.axhline(0, color="#666666", linewidth=0.7)
    ax.set_xticks(x, labels)
    ax.set_ylabel("R$ milhões")
    ax.set_title("Desempenho consolidado: 2024 × 2025")
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.6, alpha=0.75)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, ncol=2, loc="upper right")
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    fig.tight_layout()
    performance = FIG / "desempenho_2024_2025.png"
    fig.savefig(performance, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Composição do imobilizado.
    ppe = csv_rows("saidas/conciliacao_imobilizado.csv")
    names = [r["asset_class"] for r in ppe if r["asset_class"] != "TOTAL"]
    values = [int(r["net_book_value_brl"]) / 1e6 for r in ppe if r["asset_class"] != "TOTAL"]
    order = np.argsort(values)
    names = [names[i] for i in order]
    values = [values[i] for i in order]
    fig, ax = plt.subplots(figsize=(9.3, 4.3), dpi=180)
    colors = ["#D9EAF7", "#B4C7E7", "#8EA9DB", "#5B9BD5", "#4472C4", "#2F5597", "#17365D"]
    bars = ax.barh(names, values, color=colors[: len(values)])
    ax.set_xlabel("R$ milhões")
    ax.set_title("Imobilizado líquido por classe — conciliação à Nota 13")
    ax.grid(axis="x", color="#D9D9D9", linewidth=0.6, alpha=0.75)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    for bar, value in zip(bars, values):
        ax.text(value + max(values) * 0.012, bar.get_y() + bar.get_height() / 2, f"{value:,.1f}".replace(",", "X").replace(".", ",").replace("X", "."), va="center", fontsize=7.5)
    fig.tight_layout()
    ppe_path = FIG / "composicao_imobilizado.png"
    fig.savefig(ppe_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Matriz de riscos acadêmica.
    risk_rows = csv_rows("saidas/matriz_riscos.csv")
    fig, ax = plt.subplots(figsize=(8.2, 5.4), dpi=180)
    for impact in range(1, 6):
        for likelihood in range(1, 6):
            score = impact * likelihood
            color = "#E2F0D9" if score <= 6 else "#FFF2CC" if score <= 12 else "#F4B183" if score <= 19 else "#F8CBAD"
            ax.add_patch(plt.Rectangle((likelihood - 0.5, impact - 0.5), 1, 1, facecolor=color, edgecolor="white"))
    for row in risk_rows:
        xval = int(row["likelihood"])
        yval = int(row["impact"])
        ax.scatter([xval], [yval], s=250 if row["significant_risk"] == "SIM" else 150, c=NAVY.replace("#", "") if False else "#17365D", edgecolors="white", linewidth=1.0, zorder=3)
        ax.text(xval, yval, row["risk_id"], color="white", ha="center", va="center", fontsize=7, fontweight="bold", zorder=4)
    ax.set_xlim(0.5, 5.5)
    ax.set_ylim(0.5, 5.5)
    ax.set_xticks(range(1, 6))
    ax.set_yticks(range(1, 6))
    ax.set_xlabel("Probabilidade")
    ax.set_ylabel("Impacto")
    ax.set_title("Matriz acadêmica de riscos inerentes")
    ax.set_aspect("equal")
    for spine in ax.spines.values():
        spine.set_visible(False)
    fig.tight_layout()
    risk_path = FIG / "matriz_riscos.png"
    fig.savefig(risk_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    return {"performance": performance, "ppe": ppe_path, "risk": risk_path}


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.65) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(style="Legenda")
    cp.add_run(caption)


def markdown_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    parsed = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        parsed.append(cells)
    headers = parsed[0]
    rows = [row for row in parsed[2:] if row and not all(re.fullmatch(r":?-{3,}:?", x) for x in row)]
    return headers, rows


def add_markdown(doc: Document, path: Path, *, strip_first_heading: bool = False,
                 top_heading_page_break: bool = False, heading_shift: int = 0) -> None:
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    i = 0
    first_heading_seen = False
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                set_paragraph_shading(p, PALE)
                p.paragraph_format.left_indent = Inches(0.15)
                p.paragraph_format.right_indent = Inches(0.15)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, name="Aptos Mono", size=7.5, color=BLACK)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            headers, rows = markdown_table(block)
            n = len(headers)
            if n == 2:
                widths = [1.55, 4.75]
            elif n == 3:
                widths = [1.25, 2.25, 2.8]
            elif n == 4:
                widths = [0.65, 1.05, 3.6, 1.1]
            elif n == 5:
                widths = [0.55, 1.05, 2.3, 1.65, 1.1]
            else:
                widths = [6.4 / n] * n
            fs = 7.8 if n <= 3 else 6.9 if n <= 5 else 6.2
            add_dataframe_table(doc, headers, rows, widths=widths, font_size=fs)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if not first_heading_seen:
                first_heading_seen = True
                if strip_first_heading:
                    i += 1
                    continue
            effective = min(4, max(1, level + heading_shift))
            if level == 1 and top_heading_page_break:
                doc.add_page_break()
            p = doc.add_paragraph(style=f"Heading {effective}")
            add_inline(p, title, size={1: 17, 2: 13.5, 3: 11, 4: 9.5}[effective], color=NAVY if effective == 1 else BLUE if effective == 2 else DARK_GRAY)
            if effective == 1:
                set_paragraph_border(p, color=NAVY, size=14, position="bottom")
            i += 1
            continue
        numbered = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Inches(0.27)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            add_inline(p, numbered.group(2), size=9.3)
        elif bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.27)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            add_inline(p, bullet.group(1), size=9.3)
        else:
            # Junta linhas contínuas de um mesmo parágrafo Markdown.
            paragraph_lines = [line.strip()]
            j = i + 1
            while j < len(lines):
                nxt = lines[j].rstrip()
                if not nxt.strip() or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("```") or re.match(r"^\s*([-*]|\d+\.)\s+", nxt):
                    break
                paragraph_lines.append(nxt.strip())
                j += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, " ".join(paragraph_lines), size=9.3)
            i = j - 1
        i += 1


def add_toc(doc: Document) -> None:
    add_section_title(doc, "Sumário", "Mapa das seções e dos papéis de trabalho", page_break=False)
    left = [
        ("Resumo executivo", "3"),
        ("Controle documental e escopo", "4"),
        ("Entendimento da entidade e ambiente", "5"),
        ("Materialidade, estratégia e riscos", "6"),
        ("Ciclos de auditoria priorizados", "8"),
        ("Base sintética e testes", "9"),
        ("Erros inseridos e testes de 100%", "10"),
        ("Continuidade e eventos posteriores", "11"),
        ("Conclusão integrada", "12"),
        ("Papéis de trabalho", "13"),
        ("PT-EC — Estrutura Conceitual", "14"),
        ("PT-200 — Objetivos gerais", "16"),
        ("PT-210 — Termos do trabalho", "18"),
        ("PT-220 — Gestão de qualidade", "20"),
        ("PT-230 — Documentação", "22"),
        ("PT-240 — Fraude", "24"),
        ("PT-250 — Leis e regulamentos", "26"),
        ("PT-260 — Comunicação à governança", "28"),
        ("PT-265 — Deficiências de controle", "30"),
        ("PT-300 — Planejamento", "32"),
        ("PT-315 — Avaliação de riscos", "34"),
        ("PT-320 — Materialidade", "36"),
        ("PT-330 — Respostas aos riscos", "38"),
        ("PT-402 — Prestadores de serviços", "40"),
        ("PT-450 — Distorções", "42"),
    ]
    right = [
        ("PT-500 — Evidência", "44"),
        ("PT-501 — Itens selecionados", "46"),
        ("PT-505 — Confirmações", "48"),
        ("PT-510 — Saldos iniciais", "50"),
        ("PT-520 — Analíticos", "52"),
        ("PT-530 — Amostragem", "54"),
        ("PT-540 — Estimativas", "56"),
        ("PT-550 — Partes relacionadas", "58"),
        ("PT-560 — Eventos subsequentes", "60"),
        ("PT-570 — Continuidade", "62"),
        ("PT-580 — Representações", "64"),
        ("PT-600 — Auditoria de grupo", "66"),
        ("Apêndice A — Cobertura NBC TA", "68"),
        ("Apêndice B — Programas por ciclo", "70"),
        ("Apêndice C — Versões das normas", "73"),
        ("Apêndice D — Metodologia", "75"),
        ("Apêndice E — Dicionário", "77"),
        ("Apêndice F — Limitações", "79"),
        ("Apêndice G — Carta de contratação", "80"),
        ("Apêndice H — Confirmações", "81"),
        ("Apêndice I — Representação formal", "82"),
        ("Apêndice J — Componentes", "83"),
        ("Apêndice K — Referências cruzadas", "84"),
        ("Apêndice L — Qualidade", "85"),
        ("Apêndice M/N e encerramento", "86"),
    ]
    rows = [[left[i][0], left[i][1], right[i][0], right[i][1]] for i in range(25)]
    add_dataframe_table(
        doc,
        ["Seção", "Pág.", "Seção", "Pág."],
        rows,
        widths=[2.72, 0.42, 2.72, 0.42],
        font_size=6.7,
        alternate=True,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Os títulos do PDF também funcionam como marcadores de navegação.")
    set_run_font(r, size=7.5, color=DARK_GRAY, italic=True)
    doc.add_page_break()


def add_executive_part(doc: Document, figs: dict[str, Path]) -> None:
    add_section_title(doc, "Resumo executivo", "Escopo, leitura dos riscos e resultados da simulação", bookmark="ResumoExecutivo", bookmark_id=10)
    add_kpi_table(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "Este dossiê documenta um planejamento acadêmico de auditoria das demonstrações consolidadas da **Tupy S.A. e controladas** para o exercício de 2025. O trabalho cobre a Estrutura Conceitual e as NBC TA 200 a 600 exigidas na atividade, usando demonstrações financeiras, ITRs, Formulário de Referência, documentos de governança e outras fontes públicas.",
        size=9.6,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "O planejamento prioriza imobilizado e impairment, realização de tributos diferidos, receita, estoques, dívida e covenants, burla de controles pela administração e consolidação internacional. A materialidade global adotada é de **R$ 46,958 milhões**, baseada em 0,5% do ativo total. O benchmark de resultado não foi usado porque 2025 encerrou com prejuízo relevante e porque a companhia é intensiva em ativos.",
        size=9.6,
    )
    add_figure(doc, figs["performance"], "Figura 1 — Indicadores públicos selecionados das DFPs de 2025. Valores em R$ milhões. Fonte: Tupy S.A./CVM; cálculos acadêmicos.")
    add_note(
        doc,
        "Conclusão executiva",
        "As fontes públicas permitem planejar a auditoria e testar uma população artificial, mas não permitem concluir sobre a efetividade dos controles, obter confirmações externas, observar inventários, reexecutar modelos internos ou emitir opinião independente.",
    )

    add_section_title(doc, "Controle documental e escopo", None, page_break=True)
    add_dataframe_table(
        doc,
        ["Campo", "Definição"],
        [
            ["Entidade", "Tupy S.A. e controladas"],
            ["Período", "Exercício findo em 31/12/2025"],
            ["Data-base da pesquisa", "24/10/2026"],
            ["Preparador", "Murilo Sena Baiero — matrícula 261021423"],
            ["Revisor", "Em branco; não foi fabricada revisão"],
            ["Fontes", "CVM, B3, Relações com Investidores, CFC e material de disciplina"],
            ["Unidade de apresentação", "Valores em reais, salvo indicação; as DFPs públicas apresentam saldos em R$ mil"],
            ["Limite", "Trabalho acadêmico; sem acesso interno e sem opinião de auditoria"],
        ],
        widths=[1.65, 4.75],
        font_size=8.3,
    )
    add_note(
        doc,
        "Regra de classificação",
        "Cada procedimento foi identificado como **Realizado com evidência pública**, **Planejado por falta de acesso** ou **Simulado para fins didáticos**. A classificação acompanha o procedimento; ela não é inferida apenas pela seção em que aparece.",
    )

    add_section_title(doc, "Entendimento da entidade e ambiente", None, page_break=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "A Tupy atua em soluções de engenharia e manufatura, com operações industriais relevantes no Brasil, México e Portugal e estruturas societárias em outras jurisdições. A combinação de grande intensidade de ativos, estimativas complexas, moedas distintas, estoques em múltiplas plantas, dívida contratual e processos de consolidação orienta a avaliação dos riscos. A ação TUPY3 integra o Novo Mercado da B3.",
        size=9.6,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "O relatório do auditor independente sobre 2025 foi emitido pela KPMG em 19/03/2026, sem modificação de opinião, e destacou como principais assuntos a realização de ativos fiscais diferidos e o valor recuperável do imobilizado da unidade de Betim. Esses pontos foram utilizados como insumos públicos para o planejamento; a conclusão da KPMG não foi adotada como conclusão própria do aluno.",
        size=9.6,
    )
    add_dataframe_table(
        doc,
        ["Indicador", "2025", "2024", "Leitura acadêmica"],
        [
            ["Receita", "R$ 9,693 bi", "R$ 10,665 bi", "Redução de 9,1%"],
            ["Margem bruta", "12,6%", "18,1%", "Contração de 5,4 p.p."],
            ["Margem operacional", "(1,6%)", "5,2%", "Mudança para resultado negativo"],
            ["Resultado líquido", "R$ (654,6) mi", "R$ 82,4 mi", "Volatilidade e pressão patrimonial"],
            ["Dívida líquida / EBITDA ajustado", "3,35×", "1,81×", "Aumento de alavancagem"],
            ["Liquidez corrente", "2,40×", "2,20×", "Fator mitigante de curto prazo"],
        ],
        widths=[2.0, 1.15, 1.15, 2.15],
        font_size=8.0,
    )

    add_section_title(doc, "Materialidade, estratégia e riscos", None, page_break=True)
    add_dataframe_table(
        doc,
        ["Parâmetro", "Base", "%", "Valor", "Fundamentação"],
        [
            ["Materialidade global", "Ativo total", "0,5%", "R$ 46.958.205", "Benchmark estável para entidade industrial com prejuízo"],
            ["Materialidade de execução", "Materialidade global", "70%", "R$ 32.870.744", "Riscos significativos e dispersão internacional"],
            ["Claramente trivial", "Materialidade global", "5%", "R$ 2.347.910", "Limiar de acumulação de distorções"],
        ],
        widths=[1.4, 1.25, 0.65, 1.2, 2.0],
        font_size=7.8,
    )
    add_note(
        doc,
        "Julgamento profissional",
        "Os percentuais são premissas acadêmicas coerentes, e não políticas divulgadas pela KPMG ou pela Tupy. Em auditoria real, materialidade deve ser revisada quando resultados e riscos mudarem.",
        color=PALE,
    )
    add_figure(doc, figs["risk"], "Figura 2 — Matriz acadêmica R01–R12. O tamanho maior identifica risco classificado como significativo. Fonte: elaboração própria.", width=5.9)

    risks = csv_rows("saidas/matriz_riscos.csv")
    risk_table = []
    for r in risks:
        risk_table.append([
            r["risk_id"],
            r["risk"],
            r["risk_level"],
            r["significant_risk"],
            r["critical_assertions"],
            r["audit_response"],
        ])
    add_dataframe_table(
        doc,
        ["ID", "Risco", "Nível", "Signif.", "Assertivas", "Resposta planejada"],
        risk_table,
        widths=[0.38, 1.45, 0.55, 0.48, 1.25, 2.20],
        font_size=6.5,
    )

    add_section_title(doc, "Ciclos de auditoria priorizados", None, page_break=True)
    cycles = csv_rows("saidas/ciclos_auditoria.csv")
    cycle_table = [[r["cycle_id"], r["cycle"], r["linked_risks"], r["critical_assertions"], r["priority"]] for r in cycles]
    add_dataframe_table(
        doc,
        ["ID", "Ciclo", "Riscos", "Assertivas críticas", "Prioridade"],
        cycle_table,
        widths=[0.42, 2.1, 0.72, 2.45, 0.72],
        font_size=7.0,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "Os ciclos orientam os programas substantivos e de controles. Controles-chave foram identificados para teste, mas sua existência e efetividade operacional não foram confirmadas, porque walkthroughs, observação e inspeção de evidência interna não estavam disponíveis.",
        size=9.4,
    )


def add_simulation_part(doc: Document, figs: dict[str, Path]) -> None:
    add_section_title(doc, "Base sintética e testes de auditoria", "População artificial do imobilizado — 520.000 registros", page_break=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "O imobilizado foi selecionado porque é a maior conta individual do ativo em 31/12/2025, totalizando **R$ 2,525 bilhões**. A base foi construída com 520.000 registros artificiais, semente determinística vinculada à matrícula, classes, entidades, moedas, documentos, datas, autorizações, vida útil, depreciação e sinais de impairment. Apenas os totais agregados foram reconciliados à Nota 13; nenhum registro individual corresponde ao cadastro real da companhia.",
        size=9.6,
    )
    add_figure(doc, figs["ppe"], "Figura 3 — Composição da população sintética conciliada aos saldos públicos da Nota 13. Fonte: Tupy S.A./CVM e simulação própria.")

    ppe = csv_rows("saidas/conciliacao_imobilizado.csv")
    rows = []
    for r in ppe:
        rows.append([
            r["asset_class"],
            brl(int(r["book_cost_brl"]) / 1e6, 1) + " mi",
            brl(int(r["accumulated_depreciation_brl"]) / 1e6, 1) + " mi",
            brl(int(r["net_book_value_brl"]) / 1e6, 1) + " mi",
            brl(int(r["difference_net_brl"]), 0),
        ])
    add_dataframe_table(
        doc,
        ["Classe", "Custo", "Depreciação acum.", "Líquido", "Diferença"],
        rows,
        widths=[2.05, 1.1, 1.35, 1.1, 0.85],
        font_size=7.2,
    )
    add_note(
        doc,
        "Conciliação",
        "Custo de R$ 9,174 bilhões, depreciação acumulada de R$ 6,649 bilhões e saldo líquido de R$ 2,525 bilhões: diferenças iguais a zero contra os totais públicos.",
    )

    add_section_title(doc, "Erros inseridos e testes de 100%", None, page_break=True)
    tests = csv_rows("saidas/resultados_testes_100_percento.csv")
    rows = []
    for r in tests:
        amount = int(float(r["detected_misstatement_brl"]))
        rows.append([
            r["test_code"],
            r["procedure"],
            r["detected_exceptions"],
            brl(amount),
            r["test_result"],
        ])
    add_dataframe_table(
        doc,
        ["Teste", "Procedimento automatizado", "Exceções", "Distorção", "Resultado"],
        rows,
        widths=[0.48, 3.3, 0.72, 1.05, 0.82],
        font_size=7.0,
    )
    with (ROOT / "saidas" / "avaliacao_MUS.json").open("r", encoding="utf-8") as fh:
        mus = json.load(fh)
    add_dataframe_table(
        doc,
        ["Métrica MUS", "Resultado"],
        [
            ["População artificial", brl(mus["population_value_brl"])],
            ["Amostra sistemática", f"{mus['actual_unique_sample_size']} itens"],
            ["Intervalo", brl(mus["systematic_interval_brl"])],
            ["Itens com erro", str(mus["sample_error_items"])],
            ["Distorção absoluta projetada", brl(mus["projected_absolute_misstatement_brl"])],
            ["Limite superior de erro", brl(mus["upper_error_limit_brl"])],
            ["Distorção tolerável", brl(mus["tolerable_misstatement_brl"])],
            ["Conclusão", mus["conclusion"]],
        ],
        widths=[3.0, 3.4],
        font_size=8.2,
    )
    add_note(
        doc,
        "Interpretação da simulação",
        "A distorção líquida artificial é de R$ 53,0 milhões e supera a materialidade global. O limite superior MUS é de R$ 92,0 milhões e supera a distorção tolerável; por isso, a população artificial não é aceita. O resultado foi desenhado para treinamento e **não indica erro, fraude ou deficiência real na Tupy**.",
        color="FCE8E6",
    )

    add_section_title(doc, "Continuidade operacional e eventos posteriores", None, page_break=True)
    continuity = csv_rows("saidas/avaliacao_continuidade.csv")
    rows = [[r["classification"], r["factor"], r["amount_or_ratio"], r["period"], r["audit_relevance"]] for r in continuity]
    add_dataframe_table(
        doc,
        ["Classificação", "Fator", "Valor/índice", "Período", "Relevância"],
        rows,
        widths=[1.0, 1.55, 1.2, 0.75, 2.0],
        font_size=6.8,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "Há indicadores adversos — prejuízo, resultado operacional negativo, maior alavancagem e waivers — e fatores mitigantes — caixa, geração operacional, liquidez corrente e declaração de adimplência após os waivers. As evidências públicas não bastam para testar as projeções de caixa ou obter suporte contratual completo; a conclusão real sobre continuidade permanece pendente de procedimentos planejados no PT-570.",
        size=9.5,
    )

    add_section_title(doc, "Conclusão integrada", None, page_break=True)
    add_markdown(doc, ROOT / "docs" / "conclusao_integrada.md", strip_first_heading=True, heading_shift=0)


def add_workpapers(doc: Document) -> None:
    add_section_title(
        doc,
        "Papéis de trabalho — NBC TA 200 a 600",
        "Documentação individual, referências cruzadas e conclusões dentro dos limites acadêmicos",
        page_break=True,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "Os 27 papéis de trabalho a seguir preservam o conteúdo editável do repositório. A assinatura do preparador está identificada; o campo de revisão permanece em branco porque não houve segunda revisão real. As referências Fxx apontam para a matriz de fontes e as referências A-xx para anexos e saídas do próprio arquivo.",
        size=9.5,
    )
    pt_order = ["EC", "200", "210", "220", "230", "240", "250", "260", "265", "300", "315", "320", "330", "402", "450", "500", "501", "505", "510", "520", "530", "540", "550", "560", "570", "580", "600"]
    for idx, code in enumerate(pt_order, start=1):
        doc.add_page_break()
        add_markdown(doc, ROOT / "papeis_trabalho" / f"PT-{code}.md", strip_first_heading=False, top_heading_page_break=False, heading_shift=0)
        # Marcador discreto de sequência, útil em impressão.
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(f"Papel {idx:02d} de {len(pt_order)}")
        set_run_font(r, size=7.2, color=DARK_GRAY, italic=True)


def add_appendices(doc: Document) -> None:
    appendices = [
        ("Apêndice A — Matriz de cobertura NBC TA", "docs/matriz_cobertura_nbcta.md"),
        ("Apêndice B — Programas por ciclo", "docs/programas_ciclos.md"),
        ("Apêndice C — Versões das normas", "docs/versoes_normas.md"),
        ("Apêndice D — Metodologia da base sintética", "docs/metodologia_base_sintetica.md"),
        ("Apêndice E — Dicionário de dados", "base_simulada/dicionario_dados.md"),
        ("Apêndice F — Classificação e limitações", "docs/classificacao_procedimentos_e_limitacoes.md"),
        ("Apêndice G — Minuta da carta de contratação", "docs/minuta_carta_contratacao.md"),
        ("Apêndice H — Modelos de confirmações", "docs/modelos_confirmacoes.md"),
        ("Apêndice I — Minuta de representação formal", "docs/minuta_representacao_formal.md"),
        ("Apêndice J — Instruções aos componentes", "docs/instrucoes_componentes.md"),
        ("Apêndice K — Referências cruzadas", "docs/referencias_cruzadas.md"),
        ("Apêndice L — Checklist de qualidade", "docs/checklist_qualidade.md"),
    ]
    for title, rel in appendices:
        add_section_title(doc, title, None, page_break=True)
        add_markdown(doc, ROOT / rel, strip_first_heading=True, heading_shift=1)

    add_section_title(doc, "Apêndice M — Matriz de fontes", None, page_break=True)
    sources = csv_rows("fontes/fontes.csv")
    source_rows = []
    for r in sources:
        source_rows.append([
            r["source_id"],
            r["document"],
            r["issuer"],
            r["period_or_version"],
            r["access_date"],
            r["pages_used"],
            r["purpose"],
            r["url"],
        ])
    add_dataframe_table(
        doc,
        ["ID", "Documento", "Emissor", "Período", "Acesso", "Páginas", "Finalidade", "URL"],
        source_rows,
        widths=[0.3, 1.05, 0.7, 0.65, 0.65, 0.72, 1.2, 1.15],
        font_size=5.6,
    )

    add_section_title(doc, "Apêndice N — Reprodução e mapa do repositório", None, page_break=True)
    add_dataframe_table(
        doc,
        ["Caminho", "Conteúdo"],
        [
            ["dossie/", "Documento consolidado em DOCX e PDF"],
            ["papeis_trabalho/", "27 papéis editáveis em Markdown"],
            ["base_simulada/", "Base de 520.000 registros, gabarito, amostra, dicionário e hashes"],
            ["codigo/", "Gerador, testes, análises, planilha, dossiê e validador"],
            ["planilhas/", "Cálculos, conciliações, riscos, MUS e fontes em XLSX"],
            ["saidas/", "Relatórios estruturados produzidos pelos scripts"],
            ["fontes/", "Matriz de fontes públicas com URL, data e páginas"],
            ["docs/", "Metodologia, programas, minutas, versões e conclusão"],
        ],
        widths=[1.7, 4.7],
        font_size=8.2,
    )
    p = doc.add_paragraph()
    add_inline(p, "Execução principal:", size=9.5, color=NAVY)
    p = doc.add_paragraph()
    set_paragraph_shading(p, PALE)
    run = p.add_run("bash executar_tudo.sh")
    set_run_font(run, name="Aptos Mono", size=9, color=BLACK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        p,
        "A semente de geração é 261021423. O script confere quantidade de registros, totais agregados, detecção das exceções, amostra MUS e hashes. Dependências e instruções detalhadas estão no README do repositório.",
        size=9.5,
    )


def build() -> Path:
    figs = create_figures()
    doc = Document()
    configure_document(doc)
    set_headers_footers(doc)
    add_cover(doc)
    add_toc(doc)
    add_executive_part(doc, figs)
    add_simulation_part(doc, figs)
    add_workpapers(doc)
    add_appendices(doc)

    # Seção final de encerramento.
    add_section_title(doc, "Encerramento", None, page_break=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DOSSIÊ ACADÊMICO CONCLUÍDO")
    set_run_font(r, name="Aptos Display", size=18, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(
        p,
        "Preparado por Murilo Sena Baiero — matrícula 261021423\nExercício analisado: 2025 | Data de elaboração: 24/10/2026",
        size=10,
        color=DARK_GRAY,
    )
    add_note(
        doc,
        "Último lembrete",
        "Documento exclusivamente acadêmico. Não utilizar como relatório de auditoria, recomendação de investimento, parecer jurídico ou confirmação de fatos não publicados.",
        color=PALE,
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    output = build()
    print(f"Dossiê DOCX salvo em {output}")
